"""Выгрузка заказов в Excel для владельца магазина."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Order, OrderStatus
from .texts import DELIVERY, PAYMENT, STATUS

HEADERS = ["№", "Дата", "Статус", "Клиент", "Телефон", "Доставка", "Адрес", "Оплата",
           "Состав", "Товары, ₽", "Доставка, ₽", "Итого, ₽"]
WIDTHS = [6, 17, 22, 18, 15, 11, 34, 17, 42, 11, 11, 11]


def orders_xlsx(orders: list[Order]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6B4F3A")
        cell.alignment = Alignment(vertical="center")

    revenue = 0
    for o in orders:
        items = "; ".join(f"{i.title} × {i.qty}" for i in o.items)
        ws.append([o.id, o.created_at, STATUS[o.status], o.customer_name, o.phone, DELIVERY[o.delivery],
                   o.address, PAYMENT[o.payment], items, o.items_total / 100, o.delivery_price / 100, o.total / 100])
        if o.status != OrderStatus.CANCELLED.value:
            revenue += o.total

    last = ws.max_row
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "Выручка без отменённых", "", "", revenue / 100])
    ws.cell(ws.max_row, 9).font = Font(bold=True)
    ws.cell(ws.max_row, 12).font = Font(bold=True)

    for col, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[1].number_format = "DD.MM.YYYY HH:MM"
        for cell in row[9:12]:
            cell.number_format = '# ##0.00 "₽"'
        row[8].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if last >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

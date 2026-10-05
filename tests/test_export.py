from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app import services
from app.export import HEADERS, orders_xlsx
from app.models import Delivery, OrderStatus, Payment, Product
from app.services import Checkout


def test_excel_has_orders_and_revenue_without_cancelled(s):
    p = s.scalar(select(Product).where(Product.title.startswith("Колумбия")))   # 720 ₽
    for user in (1, 2):
        services.change_cart(s, user, p.id, 1)
        services.create_order(s, user, Checkout("Анна", "+79991234567", Delivery.PICKUP, "", Payment.CASH))
    services.set_status(s, 2, OrderStatus.CANCELLED)

    ws = load_workbook(BytesIO(orders_xlsx(services.all_orders(s))))["Заказы"]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert [r[0] for r in rows[1:3]] == [1, 2]
    assert rows[2][2].startswith("❌")
    assert rows[-1][8] == "Выручка без отменённых" and rows[-1][11] == 720.0
